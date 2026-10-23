import openpyxl

wb = openpyxl.load_workbook('Project_Tracker, Math mod.xlsx')
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"Sheet: {sheetname}")
    for r in range(1, 15):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
        if any(row_vals):
            print(f"  Row {r}: {row_vals}")
